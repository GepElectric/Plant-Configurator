"""Generate ALM/PAR documents from a Word template and optionally export PDF.

Usage:
    python alm_par_doc.py <input_json_path>
"""

from __future__ import annotations

import json
import os
import re
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
COMPANY_ADDRESS = "Hrv. Mornarice 1D | 21000 Split | Hrvatska"
SECTION_STYLE_FALLBACKS = ["Style1 - Header 1", "Heading 1"]
CAPTION_STYLE_FALLBACKS = ["Caption", "Normal"]
ALM_COL_W = [Cm(1.383), Cm(2.656), Cm(3.963), Cm(8.950), Cm(2.270)]
PAR_COL_W = [Cm(2.108), Cm(5.278), Cm(9.126), Cm(1.547), Cm(1.196)]
DEFAULT_LANGUAGE = "en-US"
LANGUAGE_PACKS = {
    "en-US": {
        "category": "Alarm and parameter list",
        "doc_number_label": "Document No.:",
        "author_label": "Author:",
        "section_label": "Section:",
        "section_value": "Alarm and parameter list",
        "fb_label": "FB No.:",
        "page_prefix": "Page ",
        "page_middle": " of ",
        "cover_document_label": "Document:",
        "cover_author_label": "Author:",
        "cover_customer_label": "Customer:",
        "document_history": "DOCUMENT HISTORY",
        "attachments": "Attachments",
        "document_type": "Document type",
        "type_label": "Type:",
        "date_label": "Date",
        "history_headers": ["Date", "Revision", "Page changes", "Change Reasons"],
        "attachment_headers": ["Document name", "Date", "Page number", "Revision number"],
        "alarm_section": "{index}  {name} alarm list and fault matrix",
        "parameter_section": "{index}  {name} parameter list",
        "alarm_caption": "Table {table_no}. Table of {name} alarms with IDs, descriptions, and classes.",
        "parameter_caption": "Table {table_no}. Table of {name} parameters with IDs, descriptions, initial values, and units.",
        "alarm_headers": ["Alarm no.", "Alarm system", "Alarm name", "Alarm description", "Alarm type"],
        "parameter_headers": ["Parameter no.", "Parameter name", "Parameter description", "Initial value", "Unit"],
        "table_word": "Table ",
        "toc_heading": "Table of contents:",
    },
    "nb-NO": {
        "category": "Alarm og parameterliste",
        "doc_number_label": "Dokumentnr.:",
        "author_label": "Forfatter:",
        "section_label": "Seksjon:",
        "section_value": "Alarm og parameterliste",
        "fb_label": "FB Nr.:",
        "page_prefix": "Side ",
        "page_middle": " av ",
        "cover_document_label": "Dokument:",
        "cover_author_label": "Forfatter:",
        "cover_customer_label": "Kunde:",
        "document_history": "DOKUMENTHISTORIKK",
        "attachments": "Vedlegg",
        "document_type": "Dokumenttype",
        "type_label": "Type:",
        "date_label": "Dato",
        "history_headers": ["Dato", "Revisjon", "Sideendringer", "Endringsgrunn"],
        "attachment_headers": ["Dokumentnavn", "Dato", "Sidenummer", "Revisjonsnummer"],
        "alarm_section": "{index}  {name} alarmliste og feilmatrix",
        "parameter_section": "{index}  {name} parameterliste",
        "alarm_caption": "Tabell {table_no}. Tabell over alarmer for {name} med ID-er, beskrivelser og klasser.",
        "parameter_caption": "Tabell {table_no}. Tabell over parametere for {name} med ID-er, beskrivelser, startverdier og måleenheter.",
        "alarm_headers": ["Alarmnr.", "Alarmsystem", "Alarmnavn", "Alarmbeskrivelse", "Alarmtype"],
        "parameter_headers": ["Parameternr.", "Parameternavn", "Parameterbeskrivelse", "Startverdi", "Enhet"],
        "table_word": "Tabell ",
        "toc_heading": "Innholdsfortegnelse:",
    },
}


def _normalize_lines(text: str) -> list[str]:
    return [line.strip() for line in (text or "").replace("\r\n", "\n").split("\n") if line.strip()]


def _parse_exclusion_spec(spec: str) -> list[tuple[str, int, int | None]]:
    rules: list[tuple[str, int, int | None]] = []
    for raw_token in re.split(r"[,\n;]+", spec or ""):
        token = raw_token.strip()
        if not token:
            continue
        compact = token.replace(" ", "")

        match = re.fullmatch(r"(\d+)-(\d+)", compact)
        if match:
            start = int(match.group(1))
            end = int(match.group(2))
            rules.append(("range", min(start, end), max(start, end)))
            continue

        match = re.fullmatch(r"(>=|<=|>|<|==|=)(\d+)", compact)
        if match:
            rules.append((match.group(1), int(match.group(2)), None))
            continue

        if compact.isdigit():
            rules.append(("=", int(compact), None))
            continue

        raise ValueError(f"Invalid exclusion rule '{token}'. Use values like 400, 400-500, >400, or <=120.")

    return rules


def _is_excluded(item_id: int, rules: list[tuple[str, int, int | None]]) -> bool:
    for operator, first, second in rules:
        if operator == "range" and second is not None and first <= item_id <= second:
            return True
        if operator in ("=", "==") and item_id == first:
            return True
        if operator == ">" and item_id > first:
            return True
        if operator == ">=" and item_id >= first:
            return True
        if operator == "<" and item_id < first:
            return True
        if operator == "<=" and item_id <= first:
            return True
    return False


def _apply_exclusion_spec(items: list[dict], spec: str, plc_name: str, kind: str) -> list[dict]:
    rules = _parse_exclusion_spec(spec)
    if not rules:
        return list(items)

    filtered: list[dict] = []
    for item in items:
        raw_id = item.get("id")
        try:
            item_id = int(raw_id)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{kind} item in PLC '{plc_name}' has invalid id '{raw_id}'.") from exc
        if not _is_excluded(item_id, rules):
            filtered.append(item)
    return filtered


def _resolve_language_code(meta: dict) -> str:
    requested = (meta.get("language_code") or DEFAULT_LANGUAGE).strip()
    if requested in LANGUAGE_PACKS:
        return requested
    if requested.lower().startswith("nb"):
        return "nb-NO"
    return DEFAULT_LANGUAGE


def _labels(meta: dict) -> dict[str, str]:
    return LANGUAGE_PACKS[_resolve_language_code(meta)]


def _iter_bookmark_ids(root):
    for bookmark in root.iter(qn("w:bookmarkStart")):
        raw_id = bookmark.get(qn("w:id"))
        if raw_id is None:
            continue
        try:
            yield int(raw_id)
        except ValueError:
            continue


def _next_bookmark_id(document: Document) -> int:
    maximum = 0
    maximum = max(maximum, max(_iter_bookmark_ids(document.part.element), default=0))
    for section in document.sections:
        maximum = max(maximum, max(_iter_bookmark_ids(section.header._element), default=0))
        maximum = max(maximum, max(_iter_bookmark_ids(section.footer._element), default=0))
    return maximum + 1


def _clear_paragraph(paragraph: Paragraph) -> Paragraph:
    element = paragraph._element
    p_pr = element.find(qn("w:pPr"))
    for child in list(element):
        if child is not p_pr:
            element.remove(child)
    return paragraph


def _apply_paragraph_format(paragraph: Paragraph, align=None, space_before=None, space_after=None):
    if align is not None:
        paragraph.alignment = align
    p_format = paragraph.paragraph_format
    if space_before is not None:
        p_format.space_before = space_before
    if space_after is not None:
        p_format.space_after = space_after


def _set_paragraph_text(
    paragraph: Paragraph,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
    size: Pt | None = None,
    font_name: str | None = None,
    color: RGBColor | None = None,
    align=None,
    space_before=None,
    space_after=None,
):
    _clear_paragraph(paragraph)
    _apply_paragraph_format(paragraph, align=align, space_before=space_before, space_after=space_after)
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if size is not None:
        run.font.size = size
    if font_name:
        run.font.name = font_name
    if color is not None:
        run.font.color.rgb = color
    return paragraph


def _clear_cell(cell: _Cell) -> Paragraph:
    tc = cell._tc
    tc_pr = tc.tcPr
    for child in list(tc):
        if child is not tc_pr:
            tc.remove(child)
    tc.append(OxmlElement("w:p"))
    return cell.paragraphs[0]


def _shade_cell(cell: _Cell, fill_hex: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill_hex)
    tc_pr.append(shading)


def _set_row_height(row, height_twips: int, exact: bool = False):
    tr_pr = row._tr.get_or_add_trPr()
    for old in tr_pr.findall(qn("w:trHeight")):
        tr_pr.remove(old)
    tr_height = OxmlElement("w:trHeight")
    tr_height.set(qn("w:val"), str(height_twips))
    tr_height.set(qn("w:hRule"), "exact" if exact else "atLeast")
    tr_pr.append(tr_height)


def _set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _set_table_grid(table: Table, widths):
    table.autofit = False
    tbl = table._tbl
    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(1, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width.twips)))
        tbl_grid.append(grid_col)
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width


def _resolve_style_name(document: Document, fallbacks: list[str]) -> str | None:
    style_names = {style.name for style in document.styles}
    for name in fallbacks:
        if name in style_names:
            return name
    return None


def _add_body_paragraph(document: Document, text: str = "", *, style_name: str | None = None, **fmt) -> Paragraph:
    paragraph = document.add_paragraph()
    if style_name:
        paragraph.style = style_name
    if text:
        _set_paragraph_text(paragraph, text, **fmt)
    else:
        _apply_paragraph_format(paragraph, **fmt)
    return paragraph


def _append_bookmark(paragraph: Paragraph, bookmark_id: int, bookmark_name: str):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _append_instr_field(paragraph: Paragraph, instruction: str, result_text: str, *, font_name="Arial", size=Pt(10), no_proof=False):
    begin = paragraph.add_run()
    if font_name:
        begin.font.name = font_name
    begin.font.size = size
    if no_proof:
        begin.font.no_proof = True
    begin._element.append(OxmlElement("w:fldChar"))
    begin._element[-1].set(qn("w:fldCharType"), "begin")

    instr = paragraph.add_run()
    if font_name:
        instr.font.name = font_name
    instr.font.size = size
    if no_proof:
        instr.font.no_proof = True
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    instr._element.append(instr_text)

    separate = paragraph.add_run()
    if font_name:
        separate.font.name = font_name
    separate.font.size = size
    if no_proof:
        separate.font.no_proof = True
    separate._element.append(OxmlElement("w:fldChar"))
    separate._element[-1].set(qn("w:fldCharType"), "separate")

    result = paragraph.add_run(result_text)
    if font_name:
        result.font.name = font_name
    result.font.size = size
    if no_proof:
        result.font.no_proof = True

    end = paragraph.add_run()
    if font_name:
        end.font.name = font_name
    end.font.size = size
    if no_proof:
        end.font.no_proof = True
    end._element.append(OxmlElement("w:fldChar"))
    end._element[-1].set(qn("w:fldCharType"), "end")


def _set_content_control_text(root, alias_to_text: dict[str, str]):
    for sdt in root.xpath(".//w:sdt"):
        sdt_pr = sdt.find(qn("w:sdtPr"))
        if sdt_pr is None:
            continue
        alias_node = sdt_pr.find(qn("w:alias"))
        if alias_node is None:
            continue
        alias = alias_node.get(qn("w:val"))
        if alias not in alias_to_text:
            continue
        content = sdt.find(qn("w:sdtContent"))
        if content is None:
            continue
        text_nodes = list(content.iter(qn("w:t")))
        if not text_nodes:
            continue
        replacement = alias_to_text[alias]
        text_nodes[0].text = replacement
        for extra in text_nodes[1:]:
            extra.text = ""


def _replace_first_text_in_container(container, text: str):
    text_nodes = list(container.iter(qn("w:t")))
    if not text_nodes:
        return
    text_nodes[0].text = text


def _set_update_fields_on_open(document: Document):
    settings = document.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")


def _remove_existing_sections(document: Document):
    body = document._body._element
    anchor = document.tables[4]._tbl
    remove_mode = False
    for child in list(body):
        if child is anchor:
            remove_mode = True
            continue
        if remove_mode and child.tag != qn("w:sectPr"):
            body.remove(child)


def _add_field_only_paragraph(document: Document, instruction: str, *, style_name: str | None = None, align=WD_ALIGN_PARAGRAPH.LEFT, size=Pt(10)):
    paragraph = document.add_paragraph()
    if style_name:
        paragraph.style = style_name
    _clear_paragraph(paragraph)
    _apply_paragraph_format(paragraph, align=align, space_before=Pt(0), space_after=Pt(0))
    _append_instr_field(paragraph, instruction, "", font_name="Arial", size=size)
    return paragraph


def _append_index_page(document: Document, meta: dict):
    labels = _labels(meta)
    toc_heading = _add_body_paragraph(
        document,
        labels["toc_heading"],
        bold=True,
        size=Pt(14),
        font_name="Arial",
        space_before=Pt(0),
        space_after=Pt(8),
    )
    toc_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    _add_field_only_paragraph(
        document,
        ' TOC \\o "1-3" \\h \\z \\u ',
        align=WD_ALIGN_PARAGRAPH.LEFT,
        size=Pt(12),
    )
    _add_body_paragraph(document, "", space_before=Pt(0), space_after=Pt(8))
    table_of_figures_style = _resolve_style_name(document, ["table of figures", "TableofFigures", "Normal"])
    _add_field_only_paragraph(
        document,
        ' TOC \\h \\z \\c "Table" ',
        style_name=table_of_figures_style,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        size=Pt(10),
    )


def _populate_cover_table(document: Document, meta: dict):
    labels = _labels(meta)
    customer_lines = _normalize_lines(meta.get("customer_block", ""))
    cover_cell = document.tables[1].rows[0].cells[0]
    first = _clear_cell(cover_cell)

    lines = [
        (labels["cover_document_label"], dict(bold=True, size=Pt(10))),
        (meta["title"], dict(bold=True, size=Pt(14))),
        ("", {}),
        (labels["cover_author_label"], dict(bold=True, size=Pt(10))),
        (meta.get("author", ""), dict(size=Pt(10))),
        ("Geppert Electric d.o.o.", dict(size=Pt(9))),
        (COMPANY_ADDRESS, dict(size=Pt(9))),
        ("", {}),
        (labels["cover_customer_label"], dict(bold=True, size=Pt(10))),
    ]
    for customer_line in customer_lines:
        lines.append((customer_line, dict(size=Pt(10))))

    for index, (text, fmt) in enumerate(lines):
        paragraph = first if index == 0 else cover_cell.add_paragraph()
        _set_paragraph_text(
            paragraph,
            text,
            font_name="Arial",
            align=WD_ALIGN_PARAGRAPH.LEFT,
            space_before=Pt(0),
            space_after=Pt(0),
            **fmt,
        )


def _populate_document_type_table(document: Document, meta: dict):
    labels = _labels(meta)
    title_table = document.tables[4]
    _set_paragraph_text(title_table.cell(0, 0).paragraphs[0], labels["document_type"], bold=True, size=Pt(12), font_name="Arial", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=Pt(0), space_after=Pt(0))
    _set_paragraph_text(title_table.cell(1, 0).paragraphs[0], labels["type_label"], bold=True, size=Pt(10), font_name="Arial", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=Pt(0), space_after=Pt(0))
    _set_paragraph_text(title_table.cell(1, 1).paragraphs[0], labels["section_value"], size=Pt(10), font_name="Arial", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=Pt(0), space_after=Pt(0))
    _set_paragraph_text(title_table.cell(1, 2).paragraphs[0], labels["date_label"], size=Pt(10), font_name="Arial", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=Pt(0), space_after=Pt(0))
    date_cell = document.tables[4].cell(1, 3)
    paragraph = _clear_cell(date_cell)
    _set_paragraph_text(
        paragraph,
        meta.get("date", ""),
        font_name="Arial",
        size=Pt(10),
        align=WD_ALIGN_PARAGRAPH.LEFT,
        space_before=Pt(0),
        space_after=Pt(0),
    )


def _localize_static_tables(document: Document, meta: dict):
    labels = _labels(meta)
    history = document.tables[2]
    _set_paragraph_text(history.cell(0, 0).paragraphs[0], labels["document_history"], bold=True, size=Pt(12), font_name="Arial", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=Pt(0), space_after=Pt(0))
    for idx, text in enumerate(labels["history_headers"]):
        _set_paragraph_text(history.cell(1, idx).paragraphs[0], text, size=Pt(10), font_name="Arial", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=Pt(0), space_after=Pt(0))

    attachments = document.tables[3]
    _set_paragraph_text(attachments.cell(0, 0).paragraphs[0], labels["attachments"], bold=True, size=Pt(12), font_name="Arial", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=Pt(0), space_after=Pt(0))
    for idx, text in enumerate(labels["attachment_headers"]):
        _set_paragraph_text(attachments.cell(1, idx).paragraphs[0], text, size=Pt(10), font_name="Arial", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=Pt(0), space_after=Pt(0))


def _populate_header_table(document: Document, meta: dict):
    labels = _labels(meta)
    header_table = document.sections[0].header.tables[0]
    row0 = header_table.rows[0].cells
    row2 = header_table.rows[2].cells
    row3 = header_table.rows[3].cells

    _replace_first_text_in_container(row0[2]._tc, f"{labels['doc_number_label']} ")
    _replace_first_text_in_container(row2[1]._tc, f"{labels['author_label']} ")
    _replace_first_text_in_container(row2[2]._tc, f"{labels['section_label']} ")
    _replace_first_text_in_container(row3[1]._tc, f"{labels['fb_label']} ")

    page_paragraph = _clear_cell(row3[3])
    _apply_paragraph_format(page_paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=Pt(0), space_after=Pt(0))
    prefix = page_paragraph.add_run(labels["page_prefix"])
    prefix.font.name = "Arial"
    prefix.font.size = Pt(8)
    _append_instr_field(page_paragraph, " PAGE ", "", font_name="Arial", size=Pt(8))
    middle = page_paragraph.add_run(labels["page_middle"])
    middle.font.name = "Arial"
    middle.font.size = Pt(8)
    _append_instr_field(page_paragraph, " NUMPAGES ", "", font_name="Arial", size=Pt(8))


def _set_core_properties(document: Document, meta: dict):
    labels = _labels(meta)
    customer_lines = _normalize_lines(meta.get("customer_block", ""))
    props = document.core_properties
    props.title = meta["title"]
    props.author = meta.get("author", "")
    props.comments = meta.get("doc_number", "")
    props.category = labels["category"]
    props.content_status = meta.get("fb_number", "01")
    props.subject = customer_lines[1] if len(customer_lines) > 1 else ""
    props.keywords = COMPANY_ADDRESS


def _populate_template_bindings(document: Document, meta: dict):
    labels = _labels(meta)
    customer_lines = _normalize_lines(meta.get("customer_block", ""))
    alias_map = {
        "Comments": meta.get("doc_number", ""),
        "Title": meta["title"],
        "Author": meta.get("author", ""),
        "Category": labels["category"],
        "Status": meta.get("fb_number", "01"),
        "Manager": customer_lines[0] if customer_lines else "",
        "Subject": customer_lines[1] if len(customer_lines) > 1 else "",
    }
    _set_content_control_text(document.part.element, alias_map)
    for section in document.sections:
        _set_content_control_text(section.header._element, alias_map)
        _set_content_control_text(section.footer._element, alias_map)


def _build_table(document: Document, headers: list[str], widths, rows: list[list[str]]):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_grid(table, widths)

    header_row = table.rows[0]
    _set_repeat_table_header(header_row)
    _set_row_height(header_row, 567, exact=True)

    for col_index, header in enumerate(headers):
        cell = header_row.cells[col_index]
        _shade_cell(cell, "F2F2F2")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = _clear_cell(cell)
        _set_paragraph_text(
            paragraph,
            header,
            bold=True,
            size=Pt(10),
            font_name="Arial",
            align=WD_ALIGN_PARAGRAPH.CENTER,
            space_before=Pt(0),
            space_after=Pt(0),
        )

    for values in rows:
        row = table.add_row()
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = _clear_cell(cell)
            _set_paragraph_text(
                paragraph,
                value or "",
                size=Pt(10),
                font_name="Arial",
                align=WD_ALIGN_PARAGRAPH.LEFT,
                space_before=Pt(0),
                space_after=Pt(0),
            )

    return table


def _add_heading_paragraph(document: Document, text: str, style_name: str | None, bookmark_id: int):
    paragraph = _add_body_paragraph(
        document,
        text,
        style_name=style_name,
        bold=True,
        italic=True,
        size=Pt(16),
        font_name="Arial",
        space_before=Pt(0),
        space_after=Pt(6),
    )
    _append_bookmark(paragraph, bookmark_id, f"_TocGen{bookmark_id:06d}")
    return paragraph


def _add_caption_paragraph(document: Document, labels: dict[str, str], table_no: int, caption_suffix: str, style_name: str | None, bookmark_ids: tuple[int, int]):
    paragraph = document.add_paragraph()
    if style_name:
        paragraph.style = style_name
    _clear_paragraph(paragraph)
    _apply_paragraph_format(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=Pt(0), space_after=Pt(8))
    paragraph.add_run(labels["table_word"]).font.name = "Arial"
    paragraph.runs[-1].font.size = Pt(10)
    paragraph.runs[-1].italic = True
    _append_instr_field(paragraph, " SEQ Table \\* ARABIC ", str(table_no), font_name="Arial", size=Pt(10), no_proof=True)
    suffix_run = paragraph.add_run(caption_suffix)
    suffix_run.font.name = "Arial"
    suffix_run.font.size = Pt(10)
    suffix_run.italic = True
    _append_bookmark(paragraph, bookmark_ids[0], f"_TocGen{bookmark_ids[0]:06d}")
    _append_bookmark(paragraph, bookmark_ids[1], f"_TocGen{bookmark_ids[1]:06d}")
    return paragraph


def _append_content_sections(document: Document, plcs: list[dict], meta: dict):
    labels = _labels(meta)
    heading_style = _resolve_style_name(document, SECTION_STYLE_FALLBACKS)
    caption_style = _resolve_style_name(document, CAPTION_STYLE_FALLBACKS)
    bookmark_id = _next_bookmark_id(document)
    section_no = 1
    table_no = 1

    for plc in plcs:
        plc_name = plc["plc_name"]
        alarms = _apply_exclusion_spec(plc.get("alarms", []), plc.get("alm_exclude_spec", ""), plc_name, "ALM")
        parameters = _apply_exclusion_spec(plc.get("parameters", []), plc.get("par_exclude_spec", ""), plc_name, "PAR")

        if alarms:
            document.add_page_break()
            _add_heading_paragraph(
                document,
                labels["alarm_section"].format(index=section_no, name=plc_name),
                heading_style,
                bookmark_id,
            )
            bookmark_id += 1
            _add_caption_paragraph(
                document,
                labels,
                table_no,
                labels["alarm_caption"].format(table_no=table_no, name=plc_name)[len(f"{labels['table_word']}{table_no}") :],
                caption_style,
                (bookmark_id, bookmark_id + 1),
            )
            bookmark_id += 2
            _build_table(
                document,
                labels["alarm_headers"],
                ALM_COL_W,
                [
                    [
                        str(alm.get("id", "")),
                        alm.get("system", ""),
                        alm.get("name", ""),
                        alm.get("description", ""),
                        alm.get("type", ""),
                    ]
                    for alm in alarms
                ],
            )
            section_no += 1
            table_no += 1

        if parameters:
            document.add_page_break()
            _add_heading_paragraph(
                document,
                labels["parameter_section"].format(index=section_no, name=plc_name),
                heading_style,
                bookmark_id,
            )
            bookmark_id += 1
            _add_caption_paragraph(
                document,
                labels,
                table_no,
                labels["parameter_caption"].format(table_no=table_no, name=plc_name)[len(f"{labels['table_word']}{table_no}") :],
                caption_style,
                (bookmark_id, bookmark_id + 1),
            )
            bookmark_id += 2
            _build_table(
                document,
                labels["parameter_headers"],
                PAR_COL_W,
                [
                    [
                        str(par.get("id", "")),
                        par.get("name", ""),
                        par.get("description", ""),
                        str(par.get("value", "")),
                        par.get("unit", "-"),
                    ]
                    for par in parameters
                ],
            )
            section_no += 1
            table_no += 1


def _resolve_template_path(data: dict) -> Path:
    language_code = _resolve_language_code(data.get("doc_meta", {}))
    candidate = data.get("template_path")
    if candidate:
        template_path = Path(candidate)
        if language_code == "nb-NO":
            norwegian_candidate = template_path.with_name("alm_par_template_nb.docx")
            if norwegian_candidate.exists():
                return norwegian_candidate
        if template_path.exists():
            return template_path
    templates_dir = Path(__file__).resolve().parent.parent / "templates"
    fallback_name = "alm_par_template_nb.docx" if language_code == "nb-NO" else "alm_par_template.docx"
    fallback = templates_dir / fallback_name
    if fallback.exists():
        return fallback
    raise FileNotFoundError("ALM/PAR template DOCX not found.")


def generate_docx(data: dict, output_path: str):
    template_path = _resolve_template_path(data)
    document = Document(str(template_path))
    meta = data["doc_meta"]

    _set_core_properties(document, meta)
    _populate_template_bindings(document, meta)
    _populate_header_table(document, meta)
    _populate_cover_table(document, meta)
    _localize_static_tables(document, meta)
    _populate_document_type_table(document, meta)
    _remove_existing_sections(document)
    document.add_page_break()
    _append_index_page(document, meta)
    _append_content_sections(document, data["plcs"], meta)
    _set_update_fields_on_open(document)
    document.save(output_path)


def _load_word_automation():
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise RuntimeError(
            "PDF export requires Microsoft Word and the pywin32 Python package."
        ) from exc
    return pythoncom, win32com.client


def convert_docx_to_pdf(docx_path: str, pdf_path: str, word_loader=None):
    loader = word_loader or _load_word_automation
    try:
        pythoncom, win32_client = loader()
    except ImportError as exc:
        raise RuntimeError(
            "PDF export requires Microsoft Word and the pywin32 Python package."
        ) from exc

    input_path = str(Path(docx_path).resolve())
    output_path = str(Path(pdf_path).resolve())
    app = None
    document = None

    pythoncom.CoInitialize()
    try:
        try:
            app = win32_client.gencache.EnsureDispatch("Word.Application")
        except Exception as exc:
            raise RuntimeError("Microsoft Word is not available for PDF export.") from exc

        app.Visible = False
        app.DisplayAlerts = 0
        document = app.Documents.Open(input_path, ReadOnly=False)

        story_range = document.StoryRanges
        while story_range is not None:
            try:
                story_range.Fields.Update()
            except Exception:
                pass
            try:
                story_range = story_range.NextStoryRange
            except Exception:
                break

        for index in range(1, document.TablesOfContents.Count + 1):
            document.TablesOfContents(index).Update()
        for index in range(1, document.TablesOfFigures.Count + 1):
            document.TablesOfFigures(index).Update()

        document.Repaginate()
        document.Save()
        document.ExportAsFixedFormat(output_path, 17)
    finally:
        if document is not None:
            try:
                document.Close(False)
            except Exception:
                pass
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()


def generate_pdf(data: dict, output_path: str):
    with tempfile.TemporaryDirectory(prefix="alm_par_doc_") as temp_dir:
        temp_docx = Path(temp_dir) / "working.docx"
        generate_docx(data, str(temp_docx))
        convert_docx_to_pdf(str(temp_docx), output_path)


def main():
    if len(sys.argv) < 2:
        print("Usage: python alm_par_doc.py <input_json>", file=sys.stderr)
        sys.exit(1)

    json_path = Path(sys.argv[1])
    with json_path.open("r", encoding="utf-8-sig") as handle:
        data = json.load(handle)

    formats = set(data.get("formats", ["pdf"]))
    output_folder = Path(data["output_folder"])
    base_name = data["base_filename"]
    output_folder.mkdir(parents=True, exist_ok=True)

    want_docx = "docx" in formats
    want_pdf = "pdf" in formats
    working_docx = output_folder / f"{base_name}.docx"
    generated = []

    try:
        if want_docx or want_pdf:
            generate_docx(data, str(working_docx))
            if want_docx:
                generated.append(str(working_docx))

        if want_pdf:
            pdf_path = output_folder / f"{base_name}.pdf"
            convert_docx_to_pdf(str(working_docx), str(pdf_path))
            generated.append(str(pdf_path))

        if want_pdf and not want_docx and working_docx.exists():
            working_docx.unlink()
    except Exception as exc:
        if generated:
            for item in generated:
                print(f"Generated: {item}")
        print(str(exc), file=sys.stderr)
        sys.exit(1)

    for item in generated:
        print(f"Generated: {item}")


if __name__ == "__main__":
    main()
